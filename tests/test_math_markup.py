from __future__ import annotations

import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document

from literature_reader.math_markup import append_docx_markup, render_html_markup, split_math_markup
from literature_reader.models import Annotation, FocusPoint, PaperMap, Paragraph, ReadingCopy, SourceDocument
from literature_reader.renderers import render_docx, render_html


class MathMarkupTests(unittest.TestCase):
    def test_inline_and_display_math_are_split_and_rendered_as_mathml(self) -> None:
        text = r"正文 \(R_{i,t}=\Delta^{ad}_{i,t}C_i\) 之后是 \[\frac{a+b}{c}\]。"
        chunks = split_math_markup(text)

        self.assertEqual(len(chunks), 5)
        rendered = render_html_markup(text)
        self.assertIn('<span class="math-inline">', rendered)
        self.assertIn('<span class="math-display">', rendered)
        self.assertIn("<msubsup>", rendered)
        self.assertIn("<mfrac>", rendered)

    def test_docx_uses_editable_office_math_for_subscripts_and_fractions(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "math.docx"
            document = Document()
            paragraph = document.add_paragraph()
            append_docx_markup(paragraph, r"\(R_{i,t}=\Delta^{ad}_{i,t}C_i\) and \[\frac{a}{b}\]")
            document.save(path)

            with zipfile.ZipFile(path) as archive:
                xml = archive.read("word/document.xml").decode("utf-8")
            self.assertIn("m:oMath", xml)
            self.assertIn("m:sSubSup", xml)
            self.assertIn("m:f", xml)

    def test_renderers_put_substantive_explanation_first_and_keep_math(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = SourceDocument(
                source_path=root / "paper.txt",
                title="Formula reading example",
                paragraphs=(
                    Paragraph(
                        anchor="P001",
                        text=r"The priority score is \(R_{i,t}=\Delta^{ad}_{i,t}C_i\).",
                    ),
                ),
            )
            annotation = Annotation(
                anchor="P001",
                role="这一段把风险定义为可用于排序的量。",
                context="前文给出不确定性区间，后文会用该风险量分配有限的人工审核名额。",
                explanation=(
                    r"作者把公式 \(R_{i,t}=\Delta^{ad}_{i,t}C_i\) 当作“先看哪一个案例”的依据。"
                    r"其中 \(\Delta^{ad}_{i,t}\) 表示当前案例的不确定性宽度，\(C_i\) 表示单位偏差造成的经济代价。"
                    "因此，两者任一变大都会提高审核优先级：前者说明结果更不确定，后者说明犯错更昂贵。"
                ),
                takeaway="风险排序同时考虑不确定性和后果，而不是只看预测误差。",
                caveat="该量用于决定审核优先级，不等于对未来实际损失的精确预测。",
                focus_points=(
                    FocusPoint(
                        quote=r"R_{i,t}=\Delta^{ad}_{i,t}C_i",
                        kind="formula",
                        formula_latex=r"R_{i,t}=\Delta^{ad}_{i,t}C_i",
                        explanation=(
                            "这个乘法关系把“不确定程度”和“犯错代价”放在同一把尺子上比较。"
                            "例如不确定性宽度为 2、代价为 5 时得分是 10；若代价变成 10，得分就变成 20，"
                            "因此更值得优先安排人工审核。"
                        ),
                    ),
                ),
            )
            reading_copy = ReadingCopy(
                source=source,
                paper_map=PaperMap("Formula reading example", "问题", "主张", ("定义风险",), "示例范围"),
                annotations=(annotation,),
            )
            html_path = root / "reading.html"
            docx_path = root / "reading.docx"
            render_html(reading_copy, html_path)
            render_docx(reading_copy, docx_path)

            html = html_path.read_text(encoding="utf-8")
            self.assertLess(html.index("逐段精读"), html.index("这一段在全文中做什么"))
            self.assertIn("<msubsup>", html)
            self.assertIn("source-highlight", html)
            self.assertIn("原文与讲解一一对应", html)
            with zipfile.ZipFile(docx_path) as archive:
                xml = archive.read("word/document.xml").decode("utf-8")
            self.assertIn("逐段精读", xml)
            self.assertIn("m:sSubSup", xml)


if __name__ == "__main__":
    unittest.main()
