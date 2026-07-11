"""Local extraction for text-based PDF, DOCX, and TXT inputs.

This module deliberately avoids OCR. A document without extractable text should
fail clearly instead of producing a misleading reading copy.
"""

from __future__ import annotations

import re
from pathlib import Path

from .models import Paragraph, SourceDocument


SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt"}


def extract_document(path: str | Path) -> SourceDocument:
    source_path = Path(path).expanduser().resolve()
    if not source_path.exists():
        raise FileNotFoundError(f"Input file does not exist: {source_path}")
    if source_path.suffix.lower() not in SUPPORTED_SUFFIXES:
        supported = ", ".join(sorted(SUPPORTED_SUFFIXES))
        raise ValueError(f"Unsupported input type '{source_path.suffix}'. Supported types: {supported}.")

    suffix = source_path.suffix.lower()
    if suffix == ".pdf":
        raw_paragraphs = _extract_pdf(source_path)
    elif suffix == ".docx":
        raw_paragraphs = _extract_docx(source_path)
    else:
        raw_paragraphs = _extract_txt(source_path)

    paragraphs = _to_anchored_paragraphs(raw_paragraphs)
    if not paragraphs:
        raise ValueError(
            "No readable text was extracted. This first release supports text-based PDFs; "
            "run OCR before using a scanned PDF."
        )
    title = _derive_title(raw_paragraphs, paragraphs, source_path.stem)
    return SourceDocument(source_path=source_path, title=title, paragraphs=tuple(paragraphs))


def _extract_pdf(path: Path) -> list[tuple[str, int | None]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    extracted: list[tuple[str, int | None]] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        extracted.extend((block, page_number) for block in _split_blocks(text))
    return extracted


def _extract_docx(path: Path) -> list[tuple[str, int | None]]:
    from docx import Document

    document = Document(str(path))
    return [(paragraph.text.strip(), None) for paragraph in document.paragraphs if paragraph.text.strip()]


def _extract_txt(path: Path) -> list[tuple[str, int | None]]:
    return [(block, None) for block in _split_blocks(path.read_text(encoding="utf-8"))]


def _split_blocks(text: str) -> list[str]:
    normalised = re.sub(r"[ \t]+", " ", text.replace("\r", "\n"))
    blocks = [
        re.sub(r"\s+", " ", block).strip()
        for block in re.split(r"\n\s*\n+", normalised)
    ]
    return [block for block in blocks if len(block) > 1]


def _to_anchored_paragraphs(raw_paragraphs: list[tuple[str, int | None]]) -> list[Paragraph]:
    current_section: str | None = None
    paragraphs: list[Paragraph] = []
    for text, page_number in raw_paragraphs:
        if _looks_like_heading(text):
            current_section = text
            continue
        anchor = f"P{len(paragraphs) + 1:03d}"
        paragraphs.append(
            Paragraph(
                anchor=anchor,
                text=text,
                section=current_section,
                page_number=page_number,
            )
        )
    return paragraphs


def _looks_like_heading(text: str) -> bool:
    compact = " ".join(text.split())
    if len(compact) > 110 or compact.endswith((".", ";", ":")):
        return False
    return bool(re.match(r"^(\d+(?:\.\d+)*\.?\s+)?[A-Z][A-Za-z0-9 ,&/()\-]{2,}$", compact))


def _derive_title(
    raw_paragraphs: list[tuple[str, int | None]],
    paragraphs: list[Paragraph],
    fallback: str,
) -> str:
    if raw_paragraphs:
        first_raw = raw_paragraphs[0][0].strip()
        if _looks_like_heading(first_raw) and not re.match(r"^\d+(?:\.\d+)*\.?\s+", first_raw):
            return first_raw[:160]
    first = paragraphs[0].text if paragraphs else fallback
    first_line = first.split(". ", maxsplit=1)[0].strip()
    return first_line[:160] if first_line else fallback
