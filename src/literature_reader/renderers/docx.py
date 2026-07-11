"""DOCX renderer using a two-column table to preserve source-note correspondence."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt

from ..models import Annotation, ReadingCopy


def render_docx(reading_copy: ReadingCopy, path: str | Path) -> Path:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.core_properties.title = f"{reading_copy.source.title} — 带批注的阅读版"
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)

    document.add_heading(reading_copy.source.title, level=0)
    document.add_paragraph(f"带批注的阅读版 · 原始文件：{reading_copy.source.source_path.name}")
    document.add_heading("论文整体阅读地图", level=1)
    _add_labeled_paragraph(document, "研究问题", reading_copy.paper_map.research_question)
    _add_labeled_paragraph(document, "核心主张", reading_copy.paper_map.central_claim)
    _add_labeled_paragraph(document, "论证主线", " → ".join(reading_copy.paper_map.argument_map))
    _add_labeled_paragraph(document, "适用范围与说明", reading_copy.paper_map.scope_notes)

    document.add_heading("原文对应阅读批注", level=1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.autofit = True
    table.rows[0].cells[0].text = "原文"
    table.rows[0].cells[1].text = "阅读批注"
    annotation_by_anchor = {annotation.anchor: annotation for annotation in reading_copy.annotations}

    for paragraph in reading_copy.source.paragraphs:
        row = table.add_row()
        source_cell, note_cell = row.cells
        source_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        note_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        _write_source_cell(source_cell, paragraph.anchor, paragraph.text, paragraph.section, paragraph.page_number, annotation_by_anchor[paragraph.anchor])
        _write_note_cell(note_cell, annotation_by_anchor[paragraph.anchor])

    if reading_copy.warnings:
        document.add_heading("提示", level=1)
        for warning in reading_copy.warnings:
            document.add_paragraph(warning, style="List Bullet")

    document.save(str(destination))
    return destination


def _add_labeled_paragraph(document: Document, label: str, value: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.add_run(f"{label}: ").bold = True
    paragraph.add_run(value)


def _write_source_cell(cell, anchor: str, text: str, section: str | None, page_number: int | None, annotation: Annotation) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.add_run(f"[{anchor}] ").bold = True
    if section:
        paragraph.add_run(f"{section} · ").italic = True
    if page_number:
        paragraph.add_run(f"page {page_number}")
    cell.add_paragraph(text)
    if annotation.translation:
        translation = cell.add_paragraph()
        translation.add_run("中文翻译：").bold = True
        translation.add_run(annotation.translation)


def _write_note_cell(cell, annotation: Annotation) -> None:
    cell.paragraphs[0].add_run(f"[{annotation.anchor}] 批注").bold = True
    fields = (
        ("段落作用", annotation.role),
        ("上下文关系", annotation.context),
        ("阅读解释", annotation.explanation),
        ("阅读要点", annotation.takeaway),
        ("边界与提醒", annotation.caveat),
    )
    for label, value in fields:
        paragraph = cell.add_paragraph()
        paragraph.add_run(f"{label}: ").bold = True
        paragraph.add_run(value)
